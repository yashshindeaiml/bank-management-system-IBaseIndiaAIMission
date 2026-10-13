import io
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import bank_management as bank


class BankManagementTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.workbook_path = Path(self.temporary_directory.name) / "bank.xlsx"

    def create_workbook(self, account_rows=(), transaction_rows=()):
        workbook = Workbook()
        accounts = workbook.active
        accounts.title = "Accounts"
        accounts.append(bank.ACCOUNT_HEADERS)
        for row in account_rows:
            accounts.append(row)

        transactions = workbook.create_sheet("Transactions")
        transactions.append(bank.TRANSACTION_HEADERS)
        for row in transaction_rows:
            transactions.append(row)
        workbook.save(self.workbook_path)
        workbook.close()

    def test_missing_worksheet_does_not_modify_existing_workbook(self):
        workbook = Workbook()
        accounts = workbook.active
        accounts.title = "Accounts"
        accounts.append(bank.ACCOUNT_HEADERS)
        accounts.append(["1234567", "Keep This", "555", 42])
        workbook.save(self.workbook_path)
        workbook.close()
        original_bytes = self.workbook_path.read_bytes()

        with self.assertRaisesRegex(bank.BankDataError, "Transactions"):
            bank.load_data(self.workbook_path)

        self.assertEqual(self.workbook_path.read_bytes(), original_bytes)

    def test_invalid_account_row_reports_location_without_modifying_file(self):
        self.create_workbook(account_rows=[["not-a-number", "Test", "555", 10]])
        original_bytes = self.workbook_path.read_bytes()

        with self.assertRaisesRegex(bank.BankDataError, r"Accounts worksheet, row 2"):
            bank.load_data(self.workbook_path)

        self.assertEqual(self.workbook_path.read_bytes(), original_bytes)

    def test_invalid_balance_reports_location_without_modifying_file(self):
        self.create_workbook(account_rows=[["1000001", "Test", "555", "not-money"]])
        original_bytes = self.workbook_path.read_bytes()

        with self.assertRaisesRegex(bank.BankDataError, r"Accounts worksheet, row 2, Balance"):
            bank.load_data(self.workbook_path)

        self.assertEqual(self.workbook_path.read_bytes(), original_bytes)

    def test_account_and_transaction_round_trip_as_decimal(self):
        accounts = [
            {
                "account_number": "1000001",
                "name": "Test User",
                "phone": "555",
                "balance": Decimal("1.01"),
            }
        ]
        transactions = [
            {
                "transaction_id": 1,
                "account_number": "1000001",
                "type": "Deposit",
                "amount": Decimal("1.01"),
                "date_time": "2026-10-09 12:00:00",
                "balance_after": Decimal("1.01"),
            }
        ]

        bank.save_data(accounts, transactions, self.workbook_path)
        loaded_accounts, loaded_transactions = bank.load_data(self.workbook_path)

        self.assertEqual(loaded_accounts[0]["balance"], Decimal("1.01"))
        self.assertEqual(loaded_transactions[0]["amount"], Decimal("1.01"))
        self.assertEqual(loaded_transactions[0]["balance_after"], Decimal("1.01"))

    def test_deposit_and_withdrawal_persist_after_reload(self):
        account = {
            "account_number": "1000001",
            "name": "Test",
            "phone": "555",
            "balance": Decimal("0.00"),
        }
        accounts = [account]
        transactions = []
        original_save_data = bank.save_data

        with patch.object(
            bank,
            "save_data",
            side_effect=lambda saved_accounts, saved_transactions: original_save_data(
                saved_accounts, saved_transactions, self.workbook_path
            ),
        ):
            bank.record_transaction(
                accounts, transactions, account, "Deposit", Decimal("10.01")
            )
            bank.record_transaction(
                accounts, transactions, account, "Withdrawal", Decimal("0.01")
            )

        loaded_accounts, loaded_transactions = bank.load_data(self.workbook_path)
        self.assertEqual(loaded_accounts[0]["balance"], Decimal("10.00"))
        self.assertEqual(len(loaded_transactions), 2)
        self.assertEqual(loaded_transactions[-1]["type"], "Withdrawal")
        self.assertEqual(loaded_transactions[-1]["balance_after"], Decimal("10.00"))

    def test_failed_atomic_replace_keeps_previous_workbook(self):
        self.create_workbook(account_rows=[["1000001", "Test", "555", 10]])
        original_bytes = self.workbook_path.read_bytes()

        with patch.object(bank.os, "replace", side_effect=PermissionError("file is open")):
            with self.assertRaises(OSError):
                bank.save_data([], [], self.workbook_path)

        self.assertEqual(self.workbook_path.read_bytes(), original_bytes)
        accounts, _ = bank.load_data(self.workbook_path)
        self.assertEqual(accounts[0]["name"], "Test")

    def test_failed_transaction_save_rolls_back_memory(self):
        account = {
            "account_number": "1000001",
            "name": "Test",
            "phone": "555",
            "balance": Decimal("10.00"),
        }
        accounts = [account]
        transactions = []

        with patch.object(bank, "save_data", side_effect=OSError("file is open")):
            with patch("builtins.print"):
                bank.record_transaction(
                    accounts, transactions, account, "Deposit", Decimal("1.01")
                )

        self.assertEqual(account["balance"], Decimal("10.00"))
        self.assertEqual(transactions, [])

    def test_failed_account_creation_save_removes_new_account(self):
        accounts = []
        transactions = []

        with patch("builtins.input", side_effect=["Test User", "555"]):
            with patch.object(bank, "save_data", side_effect=OSError("file is open")):
                with patch("builtins.print"):
                    bank.create_account(accounts, transactions)

        self.assertEqual(accounts, [])

    def test_all_ten_menu_options_run_without_touching_project_workbook(self):
        original_save_data = bank.save_data
        menu_inputs = [
            "invalid",
            "1",
            "Test User",
            "555",
            "2",
            "1000001",
            "12.50",
            "3",
            "1000001",
            "2.50",
            "4",
            "1000001",
            "5",
            "1000001",
            "6",
            "7",
            "1000001",
            "8",
            "1000001",
            "9",
            "10",
        ]
        output = io.StringIO()

        with patch.object(bank, "load_data", return_value=([], [])):
            with patch.object(
                bank,
                "save_data",
                side_effect=lambda accounts, transactions: original_save_data(
                    accounts, transactions, self.workbook_path
                ),
            ):
                with patch("builtins.input", side_effect=menu_inputs):
                    with patch("sys.stdout", output):
                        bank.main()

        console_output = output.getvalue()
        for expected_text in (
            "Choose a number from 1 to 10.",
            "Account created. Account number: 1000001",
            "Deposit successful. New balance: 12.50",
            "Withdrawal successful. New balance: 10.00",
            "Balance: 10.00",
            "Statement for Test User",
            "Total accounts: 1",
            "Goodbye.",
        ):
            with self.subTest(expected_text=expected_text):
                self.assertIn(expected_text, console_output)

        accounts, transactions = bank.load_data(self.workbook_path)
        self.assertEqual(accounts[0]["balance"], Decimal("10.00"))
        self.assertEqual(len(transactions), 2)

    def test_amount_input_rejects_non_finite_values_and_rounds_half_up(self):
        for value in ("NaN", "Infinity", "-Infinity"):
            with self.subTest(value=value), patch("builtins.input", return_value=value):
                with patch("builtins.print"):
                    self.assertIsNone(bank.read_amount())

        with patch("builtins.input", return_value="1.005"):
            self.assertEqual(bank.read_amount(), Decimal("1.01"))

    def test_amount_input_enforces_maximum_transaction_amount(self):
        with patch("builtins.input", return_value="1000000"):
            self.assertEqual(bank.read_amount(), Decimal("1000000.00"))

        with patch("builtins.input", return_value="1000000.01"):
            with patch("builtins.print") as print_message:
                self.assertIsNone(bank.read_amount())
        print_message.assert_called_once_with(
            "Maximum transaction amount is Rs. 10,00,000.00."
        )


if __name__ == "__main__":
    unittest.main()