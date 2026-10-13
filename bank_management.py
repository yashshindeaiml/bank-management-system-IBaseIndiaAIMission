import os
import tempfile
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook


DATA_FILE = Path(__file__).with_name("bank_data.xlsx")
ACCOUNT_HEADERS = ["Account Number", "Customer Name", "Phone", "Balance"]
TRANSACTION_HEADERS = [
    "Transaction ID",
    "Account Number",
    "Type",
    "Amount",
    "Date and Time",
    "Balance After",
]
CENT = Decimal("0.01")
MAX_TRANSACTION_AMOUNT = Decimal("1000000.00")


class BankDataError(Exception):
    pass


def parse_account_number(value, location):
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        raise BankDataError(f"{location}: account number must be a positive whole number.")

    if not number.is_finite() or number <= 0 or number != number.to_integral_value():
        raise BankDataError(f"{location}: account number must be a positive whole number.")
    return str(int(number))


def parse_money(value, location, allow_zero):
    try:
        amount = Decimal(str(value).strip())
        if not amount.is_finite():
            raise InvalidOperation
        amount = amount.quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        raise BankDataError(f"{location}: value must be a valid finite monetary amount.")

    if amount < 0 or (not allow_zero and amount == 0):
        requirement = "non-negative" if allow_zero else "greater than zero"
        raise BankDataError(f"{location}: amount must be {requirement}.")
    return amount


def validate_headers(sheet, expected, sheet_name):
    headers = tuple(cell.value for cell in sheet[1][: len(expected)])
    if headers != tuple(expected):
        raise BankDataError(
            f"The {sheet_name} worksheet has unexpected column headings. "
            "Restore the expected headings; the workbook was not changed."
        )


def save_data(accounts, transactions, file_path=DATA_FILE):
    file_path = Path(file_path)
    workbook = Workbook()
    account_sheet = workbook.active
    account_sheet.title = "Accounts"
    account_sheet.append(ACCOUNT_HEADERS)

    for account in accounts:
        account_sheet.append(
            [
                account["account_number"],
                account["name"],
                account["phone"],
                account["balance"],
            ]
        )

    transaction_sheet = workbook.create_sheet("Transactions")
    transaction_sheet.append(TRANSACTION_HEADERS)
    for transaction in transactions:
        transaction_sheet.append(
            [
                transaction["transaction_id"],
                transaction["account_number"],
                transaction["type"],
                transaction["amount"],
                transaction["date_time"],
                transaction["balance_after"],
            ]
        )

    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{file_path.stem}.",
            suffix=".tmp.xlsx",
            dir=file_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, file_path)
    except OSError as error:
        raise OSError(f"Could not save {file_path}: {error}") from error
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass


def load_data(file_path=DATA_FILE):
    file_path = Path(file_path)
    if not file_path.exists():
        save_data([], [], file_path)

    try:
        workbook = load_workbook(file_path)
    except Exception as error:
        raise BankDataError(f"Could not read workbook '{file_path}': {error}") from error

    accounts = []
    transactions = []
    try:
        missing_sheets = [
            name for name in ("Accounts", "Transactions") if name not in workbook.sheetnames
        ]
        if missing_sheets:
            names = ", ".join(missing_sheets)
            raise BankDataError(
                f"Workbook is missing required worksheet(s): {names}. "
                "The workbook was not changed."
            )

        account_sheet = workbook["Accounts"]
        transaction_sheet = workbook["Transactions"]
        validate_headers(account_sheet, ACCOUNT_HEADERS, "Accounts")
        validate_headers(transaction_sheet, TRANSACTION_HEADERS, "Transactions")

        for row_number, row in enumerate(
            account_sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if all(value is None for value in row):
                continue
            location = f"Accounts worksheet, row {row_number}"
            account_number = parse_account_number(row[0], location)
            if account_number in {account["account_number"] for account in accounts}:
                raise BankDataError(f"{location}: duplicate account number {account_number}.")
            name = str(row[1]).strip() if row[1] is not None else ""
            if not name:
                raise BankDataError(f"{location}: customer name cannot be blank.")
            balance = parse_money(row[3], f"{location}, Balance", allow_zero=True)
            accounts.append(
                {
                    "account_number": account_number,
                    "name": name,
                    "phone": str(row[2] or ""),
                    "balance": balance,
                }
            )

        transaction_ids = set()
        for row_number, row in enumerate(
            transaction_sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if all(value is None for value in row):
                continue
            location = f"Transactions worksheet, row {row_number}"
            try:
                transaction_id_value = Decimal(str(row[0]).strip())
                if (
                    not transaction_id_value.is_finite()
                    or transaction_id_value <= 0
                    or transaction_id_value != transaction_id_value.to_integral_value()
                ):
                    raise InvalidOperation
                transaction_id = int(transaction_id_value)
            except (InvalidOperation, ValueError, TypeError):
                raise BankDataError(f"{location}: transaction ID must be a positive whole number.")
            if transaction_id in transaction_ids:
                raise BankDataError(f"{location}: duplicate transaction ID {transaction_id}.")
            transaction_ids.add(transaction_id)

            account_number = parse_account_number(row[1], location)
            if not any(account["account_number"] == account_number for account in accounts):
                raise BankDataError(f"{location}: account {account_number} does not exist.")
            transaction_type = str(row[2]).strip() if row[2] is not None else ""
            if transaction_type not in ("Deposit", "Withdrawal"):
                raise BankDataError(f"{location}: transaction type must be Deposit or Withdrawal.")
            amount = parse_money(row[3], f"{location}, Amount", allow_zero=False)
            date_time = str(row[4]).strip() if row[4] is not None else ""
            if not date_time:
                raise BankDataError(f"{location}: date and time cannot be blank.")
            balance_after = parse_money(
                row[5], f"{location}, Balance After", allow_zero=True
            )
            transactions.append(
                {
                    "transaction_id": transaction_id,
                    "account_number": account_number,
                    "type": transaction_type,
                    "amount": amount,
                    "date_time": date_time,
                    "balance_after": balance_after,
                }
            )
    finally:
        workbook.close()

    return accounts, transactions


def find_account(accounts, account_number):
    for account in accounts:
        if account["account_number"] == account_number:
            return account
    return None


def create_account(accounts, transactions):
    name = input("Customer name: ").strip()
    if not name:
        print("Name cannot be empty.")
        return

    phone = input("Phone number: ").strip()
    existing_numbers = [int(account["account_number"]) for account in accounts]
    account_number = str(max(existing_numbers, default=1000000) + 1)
    account = {
        "account_number": account_number,
        "name": name,
        "phone": phone,
        "balance": Decimal("0.00"),
    }
    accounts.append(account)
    try:
        save_data(accounts, transactions)
    except OSError as error:
        accounts.pop()
        print(f"Account was not created because Excel could not be saved: {error}")
        return
    print(f"Account created. Account number: {account_number}")


def read_amount():
    try:
        amount = Decimal(input("Amount: ").strip())
    except (InvalidOperation, ValueError):
        print("Enter a valid number.")
        return None

    if not amount.is_finite() or amount <= 0:
        print("Amount must be greater than zero.")
        return None

    try:
        amount = amount.quantize(CENT, rounding=ROUND_HALF_UP)
    except InvalidOperation:
        print("Enter a valid amount with no more than two decimal places.")
        return None
    if amount == 0:
        print("Amount must be at least 0.01.")
        return None
    if amount > MAX_TRANSACTION_AMOUNT:
        print("Maximum transaction amount is Rs. 10,00,000.00.")
        return None
    return amount


def record_transaction(accounts, transactions, account, transaction_type, amount):
    if transaction_type == "Withdrawal" and amount > account["balance"]:
        print("Insufficient balance.")
        return

    previous_balance = account["balance"]
    if transaction_type == "Deposit":
        account["balance"] = previous_balance + amount
    else:
        account["balance"] = previous_balance - amount

    transaction_id = max(
        (transaction["transaction_id"] for transaction in transactions), default=0
    ) + 1
    transaction = {
        "transaction_id": transaction_id,
        "account_number": account["account_number"],
        "type": transaction_type,
        "amount": amount,
        "date_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "balance_after": account["balance"],
    }
    transactions.append(transaction)
    try:
        save_data(accounts, transactions)
    except OSError as error:
        account["balance"] = previous_balance
        transactions.pop()
        print(f"Transaction was not saved; no changes were kept: {error}")
        return
    print(f"{transaction_type} successful. New balance: {account['balance']:.2f}")


def deposit(accounts, transactions):
    account_number = input("Account number: ").strip()
    account = find_account(accounts, account_number)
    if account is None:
        print("Account not found.")
        return

    amount = read_amount()
    if amount is not None:
        record_transaction(accounts, transactions, account, "Deposit", amount)


def withdraw(accounts, transactions):
    account_number = input("Account number: ").strip()
    account = find_account(accounts, account_number)
    if account is None:
        print("Account not found.")
        return

    amount = read_amount()
    if amount is not None:
        record_transaction(accounts, transactions, account, "Withdrawal", amount)


def view_balance(accounts):
    account_number = input("Account number: ").strip()
    account = find_account(accounts, account_number)
    if account is None:
        print("Account not found.")
        return
    print(f"{account['name']} | Balance: {account['balance']:.2f}")


def search_account(accounts):
    account_number = input("Account number: ").strip()
    account = find_account(accounts, account_number)
    if account is None:
        print("Account not found.")
        return
    print(
        f"Account: {account['account_number']} | Name: {account['name']} | "
        f"Phone: {account['phone']} | Balance: {account['balance']:.2f}"
    )


def display_accounts(accounts):
    if not accounts:
        print("No accounts have been created.")
        return

    print("Account Number | Customer Name | Phone | Balance")
    for account in accounts:
        print(
            f"{account['account_number']} | {account['name']} | "
            f"{account['phone']} | {account['balance']:.2f}"
        )


def show_transactions(accounts, transactions):
    account_number = input("Account number: ").strip()
    account = find_account(accounts, account_number)
    if account is None:
        print("Account not found.")
        return

    account_transactions = [
        transaction
        for transaction in transactions
        if transaction["account_number"] == account_number
    ]
    if not account_transactions:
        print("No transactions for this account.")
        return

    print("ID | Type | Amount | Date and Time | Balance After")
    for transaction in account_transactions:
        print(
            f"{transaction['transaction_id']} | {transaction['type']} | "
            f"{transaction['amount']:.2f} | {transaction['date_time']} | "
            f"{transaction['balance_after']:.2f}"
        )


def account_statement(accounts, transactions):
    account_number = input("Account number: ").strip()
    account = find_account(accounts, account_number)
    if account is None:
        print("Account not found.")
        return

    print(
        f"Statement for {account['name']} | Account: {account_number} | "
        f"Current balance: {account['balance']:.2f}"
    )
    matching_transactions = [
        transaction
        for transaction in transactions
        if transaction["account_number"] == account_number
    ]
    if not matching_transactions:
        print("No transactions for this account.")
        return
    print("ID | Type | Amount | Date and Time | Balance After")
    for transaction in matching_transactions:
        print(
            f"{transaction['transaction_id']} | {transaction['type']} | "
            f"{transaction['amount']:.2f} | {transaction['date_time']} | "
            f"{transaction['balance_after']:.2f}"
        )


def main():
    try:
        accounts, transactions = load_data()
    except (BankDataError, OSError) as error:
        print(f"Unable to load bank data: {error}")
        return
    actions = {
        "1": lambda: create_account(accounts, transactions),
        "2": lambda: deposit(accounts, transactions),
        "3": lambda: withdraw(accounts, transactions),
        "4": lambda: view_balance(accounts),
        "5": lambda: search_account(accounts),
        "6": lambda: display_accounts(accounts),
        "7": lambda: show_transactions(accounts, transactions),
        "8": lambda: account_statement(accounts, transactions),
    }

    while True:
        print(
            "\n========== BANK MANAGEMENT SYSTEM ==========\n"
            "1. Create account\n"
            "2. Deposit money\n"
            "3. Withdraw money\n"
            "4. View balance\n"
            "5. Search account\n"
            "6. Display all accounts\n"
            "7. Transaction history\n"
            "8. Account statement\n"
            "9. Total account count\n"
            "10. Exit"
        )
        choice = input("Choose an option: ").strip()

        if choice in actions:
            actions[choice]()
        elif choice == "9":
            print(f"Total accounts: {len(accounts)}")
        elif choice == "10":
            print("Goodbye.")
            break
        else:
            print("Choose a number from 1 to 10.")


if __name__ == "__main__":
    main()